# -*- coding: utf-8 -*-
"""按《西南财经大学本科生毕业论文（设计）撰写与印制规范》生成 paper.docx。

A4；页边距上下左右均 3cm；正文小四宋体、1.5 倍行距、首行缩进 2 字符、两端对齐；
一级标题二号华文中宋加粗居中（段前 3 行、段后 2 行）；二级小三黑体；三级小四黑体；
表题在上、图题在下，黑体加粗居中，段前段后各 0.5 行；图表内容五号、单倍行距；
中文摘要、英文摘要、正文各自从第 1 页起编，页码居中、小五号。
"""
import io
import json
import os
import re

import copy as _copy

import copy as _copy

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.paragraph import Paragraph
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.paragraph import Paragraph
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "paper.md")
OUT = os.path.join(ROOT, "paper.docx")

SONG, HEI, ZS, FS, TNR, MONO = "宋体", "黑体", "华文中宋", "华文仿宋", "Times New Roman", "Consolas"
PORTRAIT_W, LANDSCAPE_W = 15.0, 23.0

PAGES_FILE = os.path.join(ROOT, "out", "toc_pages.json")
TOC_PAGES = json.load(io.open(PAGES_FILE, encoding="utf-8")) if os.path.exists(PAGES_FILE) else {}


# ------------------------------------------------------------------ 工具
def set_run_font(run, cn=SONG, en=TNR, size=None, bold=None, color=None, underline=None):
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if underline is not None:
        run.font.underline = underline


def style_font(style, cn, en=TNR, size=12, bold=False, color=RGBColor(0, 0, 0)):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)


def setup_section(sec, landscape=False):
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(3)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.75)


def set_page_start(sec, start):
    sectPr = sec._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:start"), str(start))
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(el)
    else:
        sectPr.append(el)


def footer_page_number(sec, start=None):
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    set_run_font(run, SONG, TNR, 9)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    if start is not None:
        set_page_start(sec, start)


def footer_blank(sec):
    sec.footer.is_linked_to_previous = False
    sec.footer.paragraphs[0].text = ""


def set_cell_margins(table, left=0.04, right=0.04, top=0.01, bottom=0.01):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:%s" % k)
        el.set(qn("w:w"), str(int(v * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    spec = {"top": ("single", "12"), "bottom": ("single", "12"),
            "left": ("none", "0"), "right": ("none", "0"),
            "insideH": ("none", "0"), "insideV": ("none", "0")}
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        val, sz = spec[edge]
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def header_bottom_border(table):
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        el = OxmlElement("w:bottom")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
        tcPr.append(tcBorders)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\^\[\d+\]\^)")


def parse_inline(text):
    out = []
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            out.append((tok[2:-2], True, False, False))
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            out.append((tok[1:-1], False, True, False))
        elif re.fullmatch(r"\^\[\d+\]\^", tok):
            out.append((tok[1:-1], False, False, True))
        else:
            out.append((tok, False, False, False))
    return out


def add_runs(p, text, size=12, cn=SONG, bold_all=False, bold=None):
    if bold is not None:
        bold_all = bold
    for seg, b, code, sup in parse_inline(text):
        r = p.add_run(seg)
        if code:
            set_run_font(r, MONO, MONO, size - 1.5, bold_all or b)
        else:
            set_run_font(r, cn, TNR, size, bold_all or b)
        if sup:
            r.font.superscript = True


def em_width(text):
    return sum(1.0 if ord(c) > 127 else 0.5 for c in text)


def longest_token(text):
    """最长不可断片段宽度（em）。中文可逐字断行，按 1em 计。"""
    best = 0.0
    cur = 0.0
    for ch in text:
        if ord(ch) > 127:
            best = max(best, 1.0)
            cur = 0.0
        elif ch in " ，。；：、（）()":
            best = max(best, cur)
            cur = 0.0
        else:
            cur += 0.5
            best = max(best, cur)
    return max(best, cur)


def col_units(rows):
    ncol = max(len(r) for r in rows)
    units = []
    for ci in range(ncol):
        hdr = rows[0][ci] if ci < len(rows[0]) else ""
        tok = max([longest_token(r[ci]) if ci < len(r) else 0 for r in rows] + [0])
        units.append(max(tok + 0.35, min(em_width(hdr) + 0.3, 4.2), 1.4))
    return units


def natural_cm(rows, fs=10.5):
    ncol = max(len(r) for r in rows)
    total = 0.0
    for ci in range(ncol):
        hdr = rows[0][ci] if ci < len(rows[0]) else ""
        tok = max([longest_token(r[ci]) if ci < len(r) else 0 for r in rows] + [0])
        cmax = max([em_width(r[ci]) if ci < len(r) else 0 for r in rows] + [0])
        total += max(tok + 0.6, min(em_width(hdr) + 0.3, 9.0), min(cmax, 9.0), 1.5)
    return total * fs * 0.3528 / 10.0 + 0.16 * ncol


# ------------------------------------------------------------------ 封面 / 声明
def cover_field(doc, label, value, underline=True):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(3.4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    r1 = p.add_run(label)
    set_run_font(r1, FS, TNR, 16, True)
    r2 = p.add_run(value if value else "＿" * 10)
    set_run_font(r2, FS, TNR, 16, False, underline=underline)


def blank_line(doc, n=1):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


TEMPLATE_COVER = r"D:\Zhe\Desktop\附件6：西南财经大学本科生毕业论文（设计）撰写与印制规范.docx"
TITLE_LINES = [
    "遗传算法求解 TSP 及其",
    "2-opt 局部搜索改进",
    "——基于 TSPLIB 小规模实例的复现实验",
]


def _copy_para(src_doc, dst_doc, src_para, did):
    """深拷贝段落并把图片关系迁移到目标文档。"""
    new = _copy.deepcopy(src_para._p)
    for blip in new.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        img = src_doc.part.related_parts[rid]
        new_rid = dst_doc.part.relate_to(img, RT.IMAGE)
        blip.set(qn("r:embed"), new_rid)
    for dp in new.iter(qn("wp:docPr")):
        dp.set("id", str(did))
        did += 1
    return new, did


TEMPLATE_COVER = r"D:\Zhe\Desktop\附件6：西南财经大学本科生毕业论文（设计）撰写与印制规范.docx"
FULL_TITLE = "遗传算法求解 TSP 及其 2-opt 局部搜索改进——基于 TSPLIB 小规模实例的复现实验"


def _copy_para(src_doc, dst_doc, src_para, did):
    """深拷贝段落，并把其中的图片关系迁移到目标文档。"""
    new = _copy.deepcopy(src_para._p)
    for blip in new.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        img = src_doc.part.related_parts[rid]
        blip.set(qn("r:embed"), dst_doc.part.relate_to(img, RT.IMAGE))
    for dp in new.iter(qn("wp:docPr")):
        dp.set("id", str(did))
        did += 1
    return new, did


def build_cover(doc):
    """直接采用规范文件提供的封面（校名题字 + 校徽 + 信息栏），仅填论文题目与日期。"""
    src = Document(TEMPLATE_COVER)
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    did = 900
    inserted = []
    for i in range(114, 128):
        sp = src.paragraphs[i]
        if not sp.text.strip() and not sp._p.findall(".//" + qn("w:drawing")):
            continue
        new, did = _copy_para(src, doc, sp, did)
        sectPr.addprevious(new)
        inserted.append(Paragraph(new, doc))

    for para in inserted:
        txt = para.text.strip()
        if txt == "XX年 XX月":
            done = False
            for r in para.runs:
                if "XX" in r.text:
                    r.text = "2026 年 10 月" if not done else ""
                    done = True
            continue
        for r in para.runs:
            if "阿拉伯数字" in r.text:
                r.text = ""
            if "三号，华文仿宋，下划线" not in r.text:
                continue
            if para.text.startswith("论文题目"):
                r.text = FULL_TITLE
                r.font.size = Pt(14)
                para.paragraph_format.line_spacing = 1.0
            else:
                r.text = "＿" * 13
                r.font.underline = False


def build_declaration(doc):
    for _ in range(2):
        doc.add_paragraph()
    for t in ("西南财经大学", "本科毕业论文原创性及知识产权声明"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run(t)
        set_run_font(r, ZS, TNR, 18, True)
    texts = [
        "本人郑重声明：所呈交的毕业论文是本人在导师的指导下取得的成果，论文写作严格遵循学术规范。对本论文的研究做出重要贡献的个人和集体，均已在文中以明确方式标明。因本毕业论文引起的法律结果完全由本人承担。",
        "本毕业论文成果归西南财经大学所有。",
        "特此声明",
    ]
    for t in texts:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(t)
        set_run_font(r, SONG, TNR, 12)
    for _ in range(5):
        doc.add_paragraph()
    for t in ["毕业论文作者签名：（必须手写）", "作者专业：（手写）", "作者学号：（手写）", "2026 年 10 月 20 日"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(t + "　　　　")
        set_run_font(r, SONG, TNR, 12)


# ------------------------------------------------------------------ 目录
def scan_toc(lines):
    toc = []
    for s in lines:
        s = s.strip()
        if s.startswith("## "):
            c = s[3:].strip()
            if c in ("摘　要", "摘要", "Abstract") or c.startswith("——"):
                continue
            m = re.match(r"^(\d+)\s+(.*)$", c)
            toc.append([1, "%s．%s" % (m.group(1), m.group(2)) if m else c])
        elif s.startswith("### "):
            toc.append([2, s[4:].strip()])
    return toc


def build_toc(doc, toc):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目　录")
    set_run_font(r, ZS, TNR, 22, True)

    for level, text in toc:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.first_line_indent = Pt(0)
        pf.left_indent = Cm(0) if level == 1 else Cm(0.74)
        pf.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(text)
        set_run_font(r, SONG, TNR, 12, bold=(level == 1))
        r2 = p.add_run("\t" + str(TOC_PAGES.get(text, 0)))
        set_run_font(r2, SONG, TNR, 12, bold=(level == 1))


# ------------------------------------------------------------------ 正文
class BodyWriter:
    def __init__(self, doc):
        self.doc = doc
        self.landscape = False

    def width(self):
        return LANDSCAPE_W if self.landscape else PORTRAIT_W

    def to(self, landscape):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        setup_section(sec, landscape=landscape)
        # 分节会复制上一节的 sectPr，需删除页码起始设置以免重新从 1 编号
        for old in sec._sectPr.findall(qn("w:pgNumType")):
            sec._sectPr.remove(old)
        sec.footer.is_linked_to_previous = True
        self.landscape = landscape

    def para(self, text, indent=True, size=12):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.first_line_indent = Pt(size * 2) if indent else Pt(0)
        add_runs(p, text, size)
        return p

    def heading(self, text, level):
        p = self.doc.add_paragraph(style={1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level])
        p.paragraph_format.keep_with_next = True
        if level == 1 and text in ("参考文献", "附录", "致谢"):
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
        add_runs(p, text, {1: 22, 2: 15, 3: 12}[level], ZS if level == 1 else HEI, bold=(level == 1))

    def table(self, rows, width_cm):
        ncol = max(len(r) for r in rows)
        fs = 10.5
        tbl = self.doc.add_table(rows=len(rows), cols=ncol)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        em = fs * 0.3528 / 10.0          # 每个 em 折算的厘米数
        avail = width_cm - 0.16 * ncol   # 扣除单元格左右内边距
        base, vol = [], []
        for ci in range(ncol):
            hdr = rows[0][ci] if ci < len(rows[0]) else ""
            tok = max([longest_token(r[ci]) if ci < len(r) else 0 for r in rows] + [0])
            cmax = max([em_width(r[ci]) if ci < len(r) else 0 for r in rows] + [0])
            base.append(max(tok + 0.6, min(em_width(hdr) + 0.3, 9.0), min(cmax, 9.0), 1.5))
            vol.append(sum(em_width(r[ci]) if ci < len(r) else 0 for r in rows) + 1.0)
        text_cols = {ci for ci in range(ncol)
                     if sum(em_width(r[ci]) if ci < len(r) else 0 for r in rows) / max(1, len(rows)) > 9}
        need = sum(base) * em
        if need >= avail:                # 空间不足：按比例压缩
            widths = [b * em * (avail / need) + 0.16 for b in base]
        else:                            # 有富余：按文字量分配
            extra = (avail - need) / em
            tv = sum(vol)
            widths = [(base[i] + extra * vol[i] / tv) * em + 0.16 for i in range(ncol)]

        tblPr = tbl._tbl.tblPr
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(sum(widths) * 567)))
        tblPr.append(tblW)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        set_table_borders(tbl)
        set_cell_margins(tbl)

        for ci in range(ncol):
            for row in tbl.rows:
                row.cells[ci].width = Cm(widths[ci])

        for ri, rowv in enumerate(rows):
            for ci in range(ncol):
                txt = rowv[ci] if ci < len(rowv) else ""
                cell = tbl.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci in text_cols else WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = 1.0
                pf.first_line_indent = Pt(0)
                add_runs(p, txt, fs, SONG, bold=(ri == 0))
        header_bottom_border(tbl)

        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for row in tbl.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for row in tbl.rows[:-1]:
            for cell in row.cells:
                for pp in cell.paragraphs:
                    pp.paragraph_format.keep_with_next = True

    def figure(self, path, caption):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(os.path.join(ROOT, path), width=Cm(min(14.0, self.width() - 1.0)))
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(9)
        cp.paragraph_format.space_after = Pt(9)
        cp.paragraph_format.line_spacing = 1.5
        add_runs(cp, caption, 10.5, HEI, bold=True)

    def caption(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.keep_with_next = True
        add_runs(p, text.replace("**", ""), 10.5, HEI, bold=True)

    def note(self, s):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.line_spacing = 1.0
        pf.space_before = Pt(3)
        pf.space_after = Pt(9)
        add_runs(p, s, 9)

    def code_block(self, buf):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(12)
        pf.line_spacing = 1.0
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F5F5F5")
        pPr.append(shd)
        for k, b in enumerate(buf):
            r = p.add_run(b)
            set_run_font(r, MONO, MONO, 10.5)
            if k < len(buf) - 1:
                r.add_break()

    def render(self, lines):
        i, n = 0, len(lines)
        pending = None
        while i < n:
            s = lines[i].strip()
            if not s or s == "---":
                i += 1
                continue
            if s.startswith("# ") or s.startswith("## ——"):
                i += 1
                continue
            if s.startswith("## "):
                c = s[3:].strip()
                if c in ("摘　要", "摘要"):
                    self.render_abstract_cn(lines, i)
                    i = self.next_h2(lines, i + 1)
                    continue
                if c == "Abstract":
                    self.render_abstract_en(lines, i)
                    i = self.next_h2(lines, i + 1)
                    continue
                m = re.match(r"^(\d+)\s+(.*)$", c)
                self.heading("%s．%s" % (m.group(1), m.group(2)) if m else c, 1)
                i += 1
                continue
            if s.startswith("### "):
                self.heading(s[4:].strip(), 2)
                i += 1
                continue
            if s.startswith("```"):
                i += 1
                buf = []
                while i < n and not lines[i].strip().startswith("```"):
                    buf.append(lines[i].rstrip())
                    i += 1
                i += 1
                self.code_block(buf)
                continue
            if s.startswith("!["):
                m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
                self.figure(m.group(2), m.group(1))
                i += 1
                continue
            if s.startswith("|"):
                rows = []
                while i < n and lines[i].strip().startswith("|"):
                    r = lines[i].strip()
                    if not re.fullmatch(r"\|[\s:\-\|]+\|", r):
                        rows.append([c.strip() for c in r.strip().strip("|").split("|")])
                    i += 1
                wide = natural_cm(rows) > PORTRAIT_W * 1.45
                if wide and not self.landscape:
                    self.to(True)
                if pending:
                    self.caption(pending)
                    pending = None
                self.table(rows, self.width())
                continue
            if s.startswith("**表 ") or s.startswith("**表A") or s.startswith("**附表"):
                pending = s
                i += 1
                continue
            if s.startswith("注：") or s.startswith("注 "):
                self.note(s)
                if self.landscape:
                    self.to(False)
                i += 1
                continue
            if re.match(r"^\[\d+\]\s", s):
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pf = p.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.left_indent = Pt(0)
                pf.line_spacing = 1.5
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.keep_together = True
                add_runs(p, s, 12)
                i += 1
                continue
            if s.startswith("- "):
                while i < n and lines[i].strip().startswith("- "):
                    txt = lines[i].strip()[2:]
                    p = self.doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    pf = p.paragraph_format
                    pf.first_line_indent = Pt(0)
                    pf.left_indent = Pt(24)
                    pf.line_spacing = 1.5
                    add_runs(p, "· " + txt, 12)
                    i += 1
                continue
            if re.match(r"^\d+\.\s", s):
                while i < n and re.match(r"^\d+\.\s", lines[i].strip()):
                    txt = lines[i].strip()
                    p = self.doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    pf = p.paragraph_format
                    pf.first_line_indent = Pt(0)
                    pf.left_indent = Pt(24)
                    pf.line_spacing = 1.5
                    add_runs(p, txt, 12)
                    i += 1
                continue
            if s.startswith("@@"):
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.line_spacing = 1.5
                pf.space_before = Pt(6)
                pf.space_after = Pt(6)
                add_runs(p, s[2:].strip(), 12)
                i += 1
                continue
            self.para(s, indent=True)
            i += 1

    @staticmethod
    def next_h2(lines, i):
        while i < len(lines) and not lines[i].strip().startswith("## "):
            i += 1
        return i

    def abstract_heading(self, text):
        p = self.doc.add_paragraph(style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        add_runs(p, text, 22, ZS, bold=True)

    def render_abstract_cn(self, lines, i):
        self.abstract_heading("摘　要")
        i += 1
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                continue
            if s.startswith("## ") or s == "---":
                return
            if s.startswith("**关键词**"):
                p = self.doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(24)
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(12)
                r = p.add_run("关键词：")
                set_run_font(r, SONG, TNR, 12, True)
                r2 = p.add_run(s.replace("**关键词**：", ""))
                set_run_font(r2, SONG, TNR, 12)
                i += 1
                continue
            self.para(s, indent=True)
            i += 1

    def render_abstract_en(self, lines, i):
        self.abstract_heading("Abstract")
        i += 1
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                continue
            if s.startswith("## ") or s == "---":
                return
            if s.startswith("**Keywords**"):
                p = self.doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(24)
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(12)
                r = p.add_run("Keywords: ")
                set_run_font(r, SONG, TNR, 12, True)
                r2 = p.add_run(s.split(":", 1)[1].strip())
                set_run_font(r2, SONG, TNR, 12)
                i += 1
                continue
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.first_line_indent = Pt(24)
            pf.line_spacing = 1.5
            r = p.add_run(s)
            set_run_font(r, SONG, TNR, 12)
            i += 1


# ------------------------------------------------------------------ 主流程
def build():
    lines = io.open(SRC, encoding="utf-8").read().splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    style_font(normal, SONG, TNR, 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    style_font(h1, ZS, TNR, 22, True)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(54)
    h1.paragraph_format.space_after = Pt(36)
    h1.paragraph_format.line_spacing = 1.5
    h1.paragraph_format.first_line_indent = Pt(0)

    h2 = doc.styles["Heading 2"]
    style_font(h2, HEI, TNR, 15, False)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(18)
    h2.paragraph_format.line_spacing = 1.5
    h2.paragraph_format.first_line_indent = Pt(0)

    h3 = doc.styles["Heading 3"]
    style_font(h3, HEI, TNR, 12, False)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(18)
    h3.paragraph_format.space_after = Pt(18)
    h3.paragraph_format.line_spacing = 1.5
    h3.paragraph_format.first_line_indent = Pt(0)

    sec = doc.sections[0]
    setup_section(sec)
    footer_blank(sec)
    build_cover(doc)

    sec = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(sec); footer_blank(sec)
    build_declaration(doc)

    # 切分：中文摘要 / 英文摘要 / 正文
    i_abs = next(k for k, l in enumerate(lines) if l.strip().startswith("## 摘"))
    i_en = next(k for k, l in enumerate(lines) if l.strip() == "## Abstract")
    i_body = next(k for k, l in enumerate(lines) if l.strip().startswith("## 1 "))

    sec = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(sec); footer_page_number(sec, 1)
    BodyWriter(doc).render(lines[i_abs:i_en])

    sec = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(sec); footer_page_number(sec, 1)
    BodyWriter(doc).render(lines[i_en:i_body])

    sec = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(sec); footer_blank(sec)
    build_toc(doc, scan_toc(lines))

    sec = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(sec); footer_page_number(sec, 1)
    BodyWriter(doc).render(lines[i_body:])

    try:
        doc.save(OUT)
        print("saved", OUT, "toc entries:", len(scan_toc(lines)))
    except PermissionError:
        alt = os.path.join(ROOT, "paper_v2.docx")
        doc.save(alt)
        print("saved", alt, "(paper.docx 被占用)", "toc entries:", len(scan_toc(lines)))


if __name__ == "__main__":
    build()
