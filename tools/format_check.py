# -*- coding: utf-8 -*-
"""按西南财经大学规范逐项核对 paper.docx 的版式设置。"""
import io
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

EA = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"
doc = Document("paper.docx")
out = []


def line(s=""):
    out.append(str(s))


def ea(run):
    rPr = run._element.rPr
    if rPr is None or rPr.rFonts is None:
        return None
    return rPr.rFonts.get(EA)


line("== 页面 ==")
for i, s in enumerate(doc.sections):
    line("第%d节 A4=%s 边距(上/下/左/右)=%.2f/%.2f/%.2f/%.2f cm 页眉%.2f 页脚%.2f" % (
        i, abs(s.page_width.cm - 21) < 0.1 and abs(s.page_height.cm - 29.7) < 0.1,
        s.top_margin.cm, s.bottom_margin.cm, s.left_margin.cm, s.right_margin.cm,
        s.header_distance.cm, s.footer_distance.cm))

line("\n== 样式 ==")
for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
    st = doc.styles[name]
    line("%-10s %s %s %.1fpt bold=%s" % (name, st.font.name, st.element.rPr.rFonts.get(EA) if st.element.rPr is not None else "",
                                         st.font.size.pt if st.font.size else 0, st.font.bold))

line("\n== 正文段落抽样 ==")
cnt = 0
for p in doc.paragraphs:
    if p.style.name in ("Normal", "正文") and len(p.text.strip()) > 60:
        pf = p.paragraph_format
        r = p.runs[0]
        line("首行缩进=%.0fpt 行距=%s 字号=%.1f 字体=%s/%s 对齐=%s" % (
            pf.first_line_indent.pt if pf.first_line_indent else 0, pf.line_spacing,
            r.font.size.pt if r.font.size else 0, r.font.name, ea(r), p.alignment))
        cnt += 1
        if cnt >= 3:
            break

line("\n== 标题抽样 ==")
for p in doc.paragraphs:
    if p.style.name in ("Heading 1", "Heading 2") and p.text.strip():
        r = p.runs[0]
        line("%-6s %-30s 字号=%.1f 字体=%s/%s bold=%s 对齐=%s" % (
            p.style.name, p.text.strip()[:28], r.font.size.pt if r.font.size else 0,
            r.font.name, ea(r), r.font.bold, p.alignment))
        if p.style.name == "Heading 2":
            break

line("\n== 表格内文字 ==")
for t in doc.tables[:3]:
    c = t.rows[1].cells[0]
    r = c.paragraphs[0].runs[0] if c.paragraphs[0].runs else None
    if r is not None:
        line("%d列 表内字号=%.1f 行距=%s" % (len(t.columns), r.font.size.pt if r.font.size else 0,
                                          c.paragraphs[0].paragraph_format.line_spacing))

line("\n== 页码 ==")
for i, s in enumerate(doc.sections):
    pn = s._sectPr.find(qn("w:pgNumType"))
    ftr = "有" if (s.footer.paragraphs and s.footer.paragraphs[0].text.strip()) or s.footer.part else "无"
    xml = s.footer._element.xml if s.footer._element is not None else ""
    line("第%d节 pgNumType=%s 页脚含PAGE域=%s" % (
        i, pn.get(qn("w:start")) if pn is not None else "无", "PAGE" in xml))

line("\n== 图片 ==")
z = zipfile.ZipFile("paper.docx")
line("图片数=%d" % len([n for n in z.namelist() if n.startswith("word/media")]))

io.open("out/format_check.txt", "w", encoding="utf-8").write("\n".join(out))
print("\n".join(out))
